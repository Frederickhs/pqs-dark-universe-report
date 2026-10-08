from __future__ import annotations

import csv
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from roster_reader import (
    RosterPerson,
    read_animation_roster,
    read_line_schedule,
    read_supervisor_roster,
    remove_special_population_duplicates,
)


warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style*",
    category=UserWarning,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
MAPPING_FILE = PROJECT_ROOT / "Config" / "module_mapping.csv"

OPERATIONS_FOLDER = PROJECT_ROOT / "Input" / "Maximo" / "Operations"
ANIMATION_FOLDER = PROJECT_ROOT / "Input" / "Maximo" / "Animation"

REQUIRED_HEADERS = {
    "trainee",
    "work group",
    "location",
    "wo #",
    "description",
    "module",
    "status",
    "total tasks",
    "completed tasks",
    "% complete",
}


@dataclass(frozen=True)
class MaximoChildRecord:
    sso: str
    source_type: str
    source_file: str
    work_group: str
    location: str
    work_order: str
    parent_description: str
    scheduled_start: object
    scheduled_finish: object
    parent_status: str
    parent_percent: float | None
    source_module_name: str
    module_code: str | None
    child_status: str
    total_tasks: int | float | None
    completed_tasks: int | float | None
    child_percent: float | None
    row_number: int


@dataclass(frozen=True)
class UnknownModuleRecord:
    sso: str
    source_type: str
    source_file: str
    work_order: str
    parent_description: str
    source_module_name: str
    row_number: int


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_text(value: object) -> str:
    return clean_text(value).casefold()


def normalize_sso(value: object) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def normalize_work_order(value: object) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def to_number(value: object) -> int | float | None:
    if value is None or value == "":
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    text = clean_text(value).replace(",", "")
    if not text:
        return None

    try:
        number = float(text)
    except ValueError:
        return None

    if number.is_integer():
        return int(number)
    return number


def to_percent(value: object) -> float | None:
    number = to_number(value)
    if number is None:
        return None

    result = float(number)
    if result > 1:
        result /= 100
    return result


def newest_xlsx(folder: Path) -> Path:
    files = [
        path
        for path in folder.glob("*.xlsx")
        if not path.name.startswith("~$")
    ]

    if not files:
        raise FileNotFoundError(f"No XLSX report found in: {folder}")

    return max(files, key=lambda path: path.stat().st_mtime)


def load_module_mapping() -> dict[str, str]:
    mapping: dict[str, str] = {}

    with MAPPING_FILE.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        required = {"Source Module Name", "Module Code", "Venue", "Active"}

        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError(
                "module_mapping.csv is missing one or more required columns"
            )

        for row in reader:
            if clean_text(row.get("Active")).casefold() != "yes":
                continue

            source_name = normalize_text(row.get("Source Module Name"))
            module_code = clean_text(row.get("Module Code")).upper()

            if not source_name or not module_code:
                continue

            existing = mapping.get(source_name)
            if existing and existing != module_code:
                raise ValueError(
                    f"Conflicting module mapping for {row.get('Source Module Name')!r}"
                )

            mapping[source_name] = module_code

    return mapping


def get_report_population() -> list[RosterPerson]:
    line_schedule_people = read_line_schedule()
    animation_people = read_animation_roster()
    supervisor_people = read_supervisor_roster()

    filtered_schedule, messages = remove_special_population_duplicates(
        line_schedule_people,
        animation_people,
        supervisor_people,
    )

    errors = [message for message in messages if message.startswith("ERROR:")]
    if errors:
        raise ValueError("; ".join(errors))

    return filtered_schedule + supervisor_people + animation_people


def locate_header_row(worksheet, max_search_rows: int = 30) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_search_rows),
        start=1,
    ):
        headers = {
            normalize_text(cell.value): cell.column
            for cell in row
            if clean_text(cell.value)
        }

        if REQUIRED_HEADERS.issubset(headers):
            return row_number, headers

    raise ValueError(
        f"Could not locate Maximo headers in worksheet {worksheet.title!r}"
    )


def value_from(row: tuple, headers: dict[str, int], header: str) -> object:
    column = headers.get(header)
    if column is None or column > len(row):
        return None
    return row[column - 1]


def parse_maximo_report(
    file_path: Path,
    source_type: str,
    allowed_ssos: set[str],
    mapping: dict[str, str],
) -> tuple[list[MaximoChildRecord], list[UnknownModuleRecord], list[str]]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    worksheet.reset_dimensions()

    header_row_number, headers = locate_header_row(worksheet)

    records: list[MaximoChildRecord] = []
    unknown_modules: list[UnknownModuleRecord] = []
    warnings_found: list[str] = []

    current_sso = ""
    current_parent_seen = False
    current_work_group = ""
    current_location = ""
    current_work_order = ""
    current_description = ""
    current_scheduled_start: object = None
    current_scheduled_finish: object = None
    current_parent_status = ""
    current_parent_percent: float | None = None

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=header_row_number + 1,
            values_only=True,
        ),
        start=header_row_number + 1,
    ):
        trainee = normalize_sso(value_from(row, headers, "trainee"))
        work_group = clean_text(value_from(row, headers, "work group"))
        location = clean_text(value_from(row, headers, "location"))
        work_order = normalize_work_order(value_from(row, headers, "wo #"))
        description = clean_text(value_from(row, headers, "description"))
        module_name = clean_text(value_from(row, headers, "module"))
        status = clean_text(value_from(row, headers, "status")).upper()
        total_tasks = to_number(value_from(row, headers, "total tasks"))
        completed_tasks = to_number(value_from(row, headers, "completed tasks"))
        percent_complete = to_percent(value_from(row, headers, "% complete"))

        is_parent_row = bool(
            trainee
            or work_order
            or description
            or work_group
            or location
        ) and not module_name

        if is_parent_row:
            current_parent_seen = True

            if trainee:
                current_sso = trainee
            else:
                current_sso = ""

            current_work_group = work_group
            current_location = location
            current_work_order = work_order
            current_description = description
            current_scheduled_start = value_from(row, headers, "scheduled start")
            current_scheduled_finish = value_from(row, headers, "scheduled finish")
            current_parent_status = status
            current_parent_percent = percent_complete
            continue

        if not module_name:
            continue

        if not current_sso:
            if not current_parent_seen:
                warnings_found.append(
                    f"Child row before any parent row: {file_path.name} | "
                    f"row {row_number} | {module_name}"
                )
            continue

        if current_sso not in allowed_ssos:
            continue

        module_code = mapping.get(normalize_text(module_name))

        record = MaximoChildRecord(
            sso=current_sso,
            source_type=source_type,
            source_file=file_path.name,
            work_group=current_work_group,
            location=current_location,
            work_order=current_work_order,
            parent_description=current_description,
            scheduled_start=current_scheduled_start,
            scheduled_finish=current_scheduled_finish,
            parent_status=current_parent_status,
            parent_percent=current_parent_percent,
            source_module_name=module_name,
            module_code=module_code,
            child_status=status,
            total_tasks=total_tasks,
            completed_tasks=completed_tasks,
            child_percent=percent_complete,
            row_number=row_number,
        )
        records.append(record)

        if module_code is None:
            unknown_modules.append(
                UnknownModuleRecord(
                    sso=current_sso,
                    source_type=source_type,
                    source_file=file_path.name,
                    work_order=current_work_order,
                    parent_description=current_description,
                    source_module_name=module_name,
                    row_number=row_number,
                )
            )

    workbook.close()
    return records, unknown_modules, warnings_found


def format_number(value: int | float | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def format_percent(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.1%}"


def summarize_source(
    source_name: str,
    records: list[MaximoChildRecord],
    allowed_ssos: set[str],
) -> None:
    matched_ssos = {record.sso for record in records}
    mapped_records = [record for record in records if record.module_code]
    unknown_records = [record for record in records if not record.module_code]

    print(f"\n{source_name}")
    print("-" * 78)
    print(f"Roster population available: {len(allowed_ssos)}")
    print(f"Roster SSOs found in child records: {len(matched_ssos)}")
    print(f"Child records retained: {len(records)}")
    print(f"Mapped child records: {len(mapped_records)}")
    print(f"Unknown child records: {len(unknown_records)}")

    module_codes = sorted(
        {record.module_code for record in mapped_records if record.module_code}
    )

    for module_code in module_codes:
        module_records = [
            record for record in mapped_records if record.module_code == module_code
        ]
        module_ssos = {record.sso for record in module_records}
        print(
            f"  {module_code}: {len(module_records)} child record(s) | "
            f"{len(module_ssos)} SSO(s)"
        )


def main() -> int:
    mapping = load_module_mapping()
    people = get_report_population()

    animation_ssos = {
        person.sso for person in people if person.population == "Animation"
    }
    operations_ssos = {
        person.sso for person in people if person.population != "Animation"
    }

    operations_file = newest_xlsx(OPERATIONS_FOLDER)
    animation_file = newest_xlsx(ANIMATION_FOLDER)

    operations_records, operations_unknown, operations_warnings = parse_maximo_report(
        operations_file,
        "Operations",
        operations_ssos,
        mapping,
    )

    animation_records, animation_unknown, animation_warnings = parse_maximo_report(
        animation_file,
        "Animation",
        animation_ssos,
        mapping,
    )

    print("=" * 78)
    print("PQS DARK UNIVERSE REPORT - MAXIMO CHILD SUMMARY")
    print("=" * 78)
    print(f"Active module mappings: {len(mapping)}")
    print(f"Operations report: {operations_file}")
    print(f"Animation report: {animation_file}")

    summarize_source("OPERATIONS", operations_records, operations_ssos)
    summarize_source("ANIMATION", animation_records, animation_ssos)

    unknown_modules = operations_unknown + animation_unknown
    parser_warnings = operations_warnings + animation_warnings

    print("\nUNKNOWN MODULES")
    print("-" * 78)
    if unknown_modules:
        for record in unknown_modules:
            print(
                f"  {record.source_type} | {record.sso} | {record.work_order} | "
                f"{record.source_module_name} | row {record.row_number}"
            )
    else:
        print("  No unknown child module descriptions found for report employees.")

    print("\nPARSER WARNINGS")
    print("-" * 78)
    if parser_warnings:
        for warning in parser_warnings:
            print(f"  {warning}")
    else:
        print("  No parent-child parsing warnings found.")

    print("\nSAMPLE RETAINED CHILD RECORDS")
    print("-" * 78)
    samples = (operations_records + animation_records)[:20]

    if not samples:
        print("  No child records matched the report population.")
    else:
        for record in samples:
            progress = (
                f"{format_number(record.completed_tasks)}/"
                f"{format_number(record.total_tasks)}"
            )
            print(
                f"  {record.source_type} | {record.sso} | {record.work_order} | "
                f"{record.module_code or 'UNKNOWN'} | {record.child_status} | "
                f"{progress} | {format_percent(record.child_percent)}"
            )

    print("\nMAXIMO PARSING COMPLETED.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
