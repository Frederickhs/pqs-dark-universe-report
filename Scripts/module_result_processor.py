from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from maximo_reader import (
    ANIMATION_FOLDER,
    OPERATIONS_FOLDER,
    MaximoChildRecord,
    load_module_mapping,
    newest_xlsx,
    parse_maximo_report,
)
from qualification_reader import build_qualification_lookup
from roster_reader import RosterPerson
from qualification_reader import get_report_population


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SETTINGS_FILE = PROJECT_ROOT / "Config" / "project_settings.json"
EXCEPTIONS_FILE = PROJECT_ROOT / "Input" / "Exceptions" / "pqs_exceptions.xlsx"


@dataclass(frozen=True)
class ModuleException:
    sso: str
    module_code: str
    exception_type: str
    reason: str
    active: bool
    effective_date: object
    end_date: object


@dataclass(frozen=True)
class ModuleResult:
    sso: str
    module_code: str
    result: str
    completed_tasks: int | float | None
    total_tasks: int | float | None
    percentage: float | None
    work_order: str
    child_status: str
    source: str
    note: str
    qualification_date: object = None


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_sso(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = clean_text(value)
    if not text:
        return None

    for pattern in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def load_settings() -> dict:
    with SETTINGS_FILE.open("r", encoding="utf-8") as file:
        return json.load(file)


def required_modules_for(person: RosterPerson, settings: dict) -> list[str]:
    groups = settings["required_modules"]

    if person.population in {"1st Shift", "2nd Shift"}:
        return list(groups["dark_universe"])
    if person.population == "3rd Shift Wolf":
        return list(groups["wolf"])
    if person.population == "3rd Shift Unchained":
        return list(groups["unchained"])
    if person.population == "Animation":
        return list(groups["animation"])
    if person.population == "Supervisors":
        if person.shift in {"1st", "2nd"} or person.venue == "Dark Universe":
            return list(groups["dark_universe"])
        if person.venue == "Wolf":
            return list(groups["wolf"])
        if person.venue == "Unchained":
            return list(groups["unchained"])

    raise ValueError(
        f"Unable to assign required modules: {person.sso} | "
        f"{person.population} | {person.shift} | {person.venue}"
    )


def find_header_row(worksheet, required: set[str], max_rows: int = 20):
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_rows), start=1
    ):
        headers = {
            clean_text(cell.value).casefold(): cell.column
            for cell in row
            if clean_text(cell.value)
        }
        if required.issubset(headers):
            return row_number, headers
    raise ValueError(f"Required headers not found in {worksheet.title}")


def row_value(row: tuple, headers: dict[str, int], name: str) -> object:
    column = headers.get(name)
    if column is None or column > len(row):
        return None
    return row[column - 1]


def yes_value(value: object) -> bool:
    return clean_text(value).casefold() in {"yes", "y", "true", "1", "active"}


def read_exceptions(today: date | None = None) -> tuple[dict[tuple[str, str], ModuleException], list[str]]:
    today = today or date.today()
    workbook = load_workbook(EXCEPTIONS_FILE, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    worksheet.reset_dimensions()

    required = {"sso", "module code", "exception type", "reason", "active"}
    header_row, headers = find_header_row(worksheet, required)

    lookup: dict[tuple[str, str], ModuleException] = {}
    warnings: list[str] = []

    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        sso = normalize_sso(row_value(row, headers, "sso"))
        module_code = clean_text(row_value(row, headers, "module code")).upper()
        exception_type = clean_text(row_value(row, headers, "exception type"))
        reason = clean_text(row_value(row, headers, "reason"))
        active_flag = yes_value(row_value(row, headers, "active"))
        effective_raw = row_value(row, headers, "effective date")
        end_raw = row_value(row, headers, "end date")

        if not sso and not module_code:
            continue
        if not sso or not module_code:
            warnings.append(f"Incomplete exception row: SSO={sso!r}, module={module_code!r}")
            continue

        effective = parse_date(effective_raw)
        end = parse_date(end_raw)
        in_date_range = (effective is None or effective <= today) and (end is None or today <= end)
        is_active = active_flag and in_date_range

        record = ModuleException(
            sso=sso,
            module_code=module_code,
            exception_type=exception_type,
            reason=reason,
            active=is_active,
            effective_date=effective_raw,
            end_date=end_raw,
        )

        key = (sso, module_code)
        if key in lookup:
            warnings.append(f"Duplicate exception row: {sso} | {module_code}")
        lookup[key] = record

        if is_active and exception_type.casefold() != "not required":
            warnings.append(
                f"Unsupported active exception type: {sso} | {module_code} | {exception_type}"
            )

    workbook.close()
    return lookup, warnings


def record_date_key(record: MaximoChildRecord) -> datetime:
    for value in (record.scheduled_finish, record.scheduled_start):
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        text = clean_text(value)
        for pattern in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, pattern)
            except ValueError:
                continue
    return datetime.min


def numeric(value: int | float | None) -> float:
    return float(value) if value is not None else 0.0


def is_officially_completed(record: MaximoChildRecord) -> bool:
    return record.child_status == "COMP" and (
        (record.child_percent is not None and record.child_percent >= 0.999999)
        or (
            record.total_tasks is not None
            and record.completed_tasks is not None
            and numeric(record.completed_tasks) >= numeric(record.total_tasks)
        )
    )


def choose_maximo_record(records: Iterable[MaximoChildRecord]) -> MaximoChildRecord | None:
    records = list(records)
    if not records:
        return None

    completed = [record for record in records if is_officially_completed(record)]
    if completed:
        return max(completed, key=lambda record: (record_date_key(record), record.row_number))

    return max(
        records,
        key=lambda record: (
            numeric(record.child_percent),
            numeric(record.completed_tasks),
            record_date_key(record),
            record.row_number,
        ),
    )


def remaining_tasks(record: MaximoChildRecord) -> float | None:
    if record.total_tasks is None:
        return None
    return max(0.0, numeric(record.total_tasks) - numeric(record.completed_tasks))


def evaluate_maximo_record(record: MaximoChildRecord) -> ModuleResult:
    remaining = remaining_tasks(record)
    completed = numeric(record.completed_tasks)
    total = numeric(record.total_tasks)

    if record.child_status == "ACK":
        result = "ACK Error - Technician Action Required"
        note = f"ACK Error: {record.module_code}"
    elif is_officially_completed(record):
        result = "Completed"
        note = ""
    elif remaining == 0 and record.total_tasks is not None:
        result = "Completed - Awaiting Maximo Status"
        note = f"Awaiting Maximo Status: {record.module_code}"
    elif remaining in {1, 2}:
        result = "Review Required"
        note = f"Review Required: {record.module_code}"
    elif completed <= 0 and total > 0:
        result = "No Progress"
        note = f"No Progress: {record.module_code}"
    else:
        result = "In Progress"
        note = f"In Progress: {record.module_code}"

    return ModuleResult(
        sso=record.sso,
        module_code=record.module_code or "UNKNOWN",
        result=result,
        completed_tasks=record.completed_tasks if record.completed_tasks is not None else 0,
        total_tasks=record.total_tasks,
        percentage=record.child_percent,
        work_order=record.work_order,
        child_status=record.child_status,
        source=record.source_type,
        note=note,
    )


def build_results():
    settings = load_settings()
    people = get_report_population()
    people_by_sso = {person.sso: person for person in people}
    population_ssos = set(people_by_sso)

    qualification_lookup, qualification_warnings = build_qualification_lookup(population_ssos)
    exception_lookup, exception_warnings = read_exceptions()

    mapping = load_module_mapping()
    animation_ssos = {p.sso for p in people if p.population == "Animation"}
    operations_ssos = population_ssos - animation_ssos

    operations_records, operations_unknown, operations_parser_warnings = parse_maximo_report(
        newest_xlsx(OPERATIONS_FOLDER), "Operations", operations_ssos, mapping
    )
    animation_records, animation_unknown, animation_parser_warnings = parse_maximo_report(
        newest_xlsx(ANIMATION_FOLDER), "Animation", animation_ssos, mapping
    )

    maximo_records = operations_records + animation_records
    maximo_by_key: dict[tuple[str, str], list[MaximoChildRecord]] = {}
    maximo_ssos = {record.sso for record in maximo_records}

    for record in maximo_records:
        if record.module_code:
            maximo_by_key.setdefault((record.sso, record.module_code), []).append(record)

    results: dict[tuple[str, str], ModuleResult] = {}

    for person in people:
        for module_code in required_modules_for(person, settings):
            key = (person.sso, module_code)
            exception = exception_lookup.get(key)

            if exception and exception.active and exception.exception_type.casefold() == "not required":
                results[key] = ModuleResult(
                    sso=person.sso,
                    module_code=module_code,
                    result="Not Required",
                    completed_tasks=None,
                    total_tasks=None,
                    percentage=None,
                    work_order="",
                    child_status="",
                    source="Exception",
                    note=f"Exception: {module_code} not required - {exception.reason}".rstrip(" -"),
                )
                continue

            qualification = qualification_lookup.get(key)
            if qualification and qualification.status == "ACTIVE":
                results[key] = ModuleResult(
                    sso=person.sso,
                    module_code=module_code,
                    result="Completed",
                    completed_tasks=None,
                    total_tasks=None,
                    percentage=None,
                    work_order="",
                    child_status="ACTIVE",
                    source="Qualification Report",
                    note="",
                    qualification_date=qualification.effective_date,
                )
                continue

            selected = choose_maximo_record(maximo_by_key.get(key, []))
            if selected:
                results[key] = evaluate_maximo_record(selected)
                continue

            if person.sso in maximo_ssos:
                missing_result = "Not Started"
                note = f"Missing Module Record: {module_code}"
            else:
                missing_result = "No Maximo Record"
                note = f"No Maximo Record: {module_code}"

            results[key] = ModuleResult(
                sso=person.sso,
                module_code=module_code,
                result=missing_result,
                completed_tasks=None,
                total_tasks=None,
                percentage=None,
                work_order="",
                child_status="",
                source="",
                note=note,
            )

    warnings = (
        qualification_warnings
        + exception_warnings
        + operations_parser_warnings
        + animation_parser_warnings
    )

    unknown_modules = operations_unknown + animation_unknown
    return people, results, warnings, unknown_modules


def progress_text(result: ModuleResult) -> str:
    if result.completed_tasks is None and result.total_tasks is None:
        return ""

    completed = result.completed_tasks if result.completed_tasks is not None else 0
    total = result.total_tasks if result.total_tasks is not None else ""
    return f"{completed}/{total}"


def percentage_text(result: ModuleResult) -> str:
    if result.percentage is None:
        return ""
    return f"{result.percentage:.1%}"


def main() -> int:
    people, results, warnings, unknown_modules = build_results()
    settings = load_settings()

    counts: dict[str, int] = {}
    for result in results.values():
        counts[result.result] = counts.get(result.result, 0) + 1

    print("=" * 82)
    print("PQS DARK UNIVERSE REPORT - INITIAL MODULE RESULTS")
    print("=" * 82)
    print(f"Employees: {len(people)}")
    print(f"Required employee-module results: {len(results)}")

    print("\nRESULT COUNTS")
    print("-" * 82)
    for status, count in sorted(counts.items()):
        print(f"  {status}: {count}")

    print("\nEXCEPTION RESULTS")
    print("-" * 82)
    exception_results = [r for r in results.values() if r.source == "Exception"]
    if exception_results:
        for result in exception_results:
            person = next(p for p in people if p.sso == result.sso)
            print(f"  {result.sso} | {person.name} | {result.module_code} | {result.note}")
    else:
        print("  No active exceptions applied.")

    print("\nNON-COMPLETED SAMPLE")
    print("-" * 82)
    sample_count = 0
    for person in people:
        for module_code in required_modules_for(person, settings):
            result = results[(person.sso, module_code)]
            if result.result in {"Completed", "Not Required"}:
                continue
            print(
                f"  {person.population} | {result.sso} | {person.name} | "
                f"{module_code} | {result.result} | {progress_text(result)} | "
                f"{percentage_text(result)} | WO {result.work_order}"
            )
            sample_count += 1
            if sample_count >= 40:
                break
        if sample_count >= 40:
            break

    print("\nWARNINGS")
    print("-" * 82)
    if warnings:
        for warning in warnings:
            print(f"  {warning}")
    else:
        print("  No processing warnings found.")

    print("\nUNKNOWN MODULES")
    print("-" * 82)
    if unknown_modules:
        for record in unknown_modules:
            print(
                f"  {record.source_type} | {record.sso} | {record.work_order} | "
                f"{record.source_module_name}"
            )
    else:
        print("  No unknown module descriptions found.")

    print("\nINITIAL MODULE RESULT PROCESSING PASSED.")
    print("Weekly snapshot comparison is not applied yet.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
