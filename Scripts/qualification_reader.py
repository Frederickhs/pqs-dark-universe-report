from __future__ import annotations

import json
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style*",
    category=UserWarning,
)

from roster_reader import (
    RosterPerson,
    read_animation_roster,
    read_line_schedule,
    read_supervisor_roster,
    remove_special_population_duplicates,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SETTINGS_FILE = PROJECT_ROOT / "Config" / "project_settings.json"

QUALIFICATION_FILES = {
    "CUMOD0": "CUMOD0 Qualification.xlsx",
    "CUMOD1": "CUMOD1 Qualification.xlsx",
    "CUMOD2": "CUMOD2 Qualification.xlsx",
    "CUMOD3": "CUMOD3 Qualification.xlsx",
    "DUMOD0": "DUMOD0 Qualification.xlsx",
    "DUMOD1": "DUMOD1 Qualification.xlsx",
    "UCMOD0": "UCMOD0 Qualification.xlsx",
    "UCMOD1": "UCMOD1 Qualification.xlsx",
    "UCMOD2": "UCMOD2 Qualification.xlsx",
    "UCMOD3": "UCMOD3 Qualification.xlsx",
    "UCMOD4": "UCMOD4 Qualification.xlsx",
}


@dataclass(frozen=True)
class QualificationRecord:
    sso: str
    module_code: str
    name: str
    effective_date: date | datetime | None
    expiration_date: date | datetime | None
    status: str
    work_site: str
    work_location: str
    source_file: str


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_sso(value: object) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]

    return text


def normalize_header(value: object) -> str:
    return clean_text(value).casefold()


def load_settings() -> dict:
    with SETTINGS_FILE.open("r", encoding="utf-8") as file:
        return json.load(file)


def get_report_population() -> list[RosterPerson]:
    line_schedule_people = read_line_schedule()
    animation_people = read_animation_roster()
    supervisor_people = read_supervisor_roster()

    filtered_schedule, messages = remove_special_population_duplicates(
        line_schedule_people,
        animation_people,
        supervisor_people,
    )

    overlap_errors = [message for message in messages if message.startswith("ERROR:")]
    if overlap_errors:
        raise ValueError("; ".join(overlap_errors))

    return filtered_schedule + supervisor_people + animation_people


def locate_header_row(worksheet, max_search_rows: int = 25) -> tuple[int, dict[str, int]]:
    required_headers = {"labor", "name", "effective date", "status"}

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=max_search_rows),
        start=1,
    ):
        headers = {
            normalize_header(cell.value): cell.column
            for cell in row
            if clean_text(cell.value)
        }

        if required_headers.issubset(headers):
            return row_number, headers

    raise ValueError(
        f"Could not find qualification headers in worksheet {worksheet.title!r}"
    )


def cell_value(row: tuple, headers: dict[str, int], header_name: str) -> object:
    column_number = headers.get(header_name)
    if column_number is None or column_number > len(row):
        return None
    return row[column_number - 1]


def read_qualification_file(
    file_path: Path,
    module_code: str,
    population_ssos: set[str],
) -> list[QualificationRecord]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)

    if "List of Labor" in workbook.sheetnames:
        worksheet = workbook["List of Labor"]
    else:
        worksheet = workbook[workbook.sheetnames[0]]

    worksheet.reset_dimensions()
    header_row_number, headers = locate_header_row(worksheet)
    records: list[QualificationRecord] = []

    for row in worksheet.iter_rows(
        min_row=header_row_number + 1,
        values_only=True,
    ):
        sso = normalize_sso(cell_value(row, headers, "labor"))
        if not sso or sso not in population_ssos:
            continue

        records.append(
            QualificationRecord(
                sso=sso,
                module_code=module_code,
                name=clean_text(cell_value(row, headers, "name")),
                effective_date=cell_value(row, headers, "effective date"),
                expiration_date=cell_value(row, headers, "expiration date"),
                status=clean_text(cell_value(row, headers, "status")).upper(),
                work_site=clean_text(cell_value(row, headers, "work site")),
                work_location=clean_text(cell_value(row, headers, "work location")),
                source_file=file_path.name,
            )
        )

    workbook.close()
    return records


def choose_record(records: list[QualificationRecord]) -> QualificationRecord:
    def sort_key(record: QualificationRecord) -> tuple[int, datetime]:
        active_rank = 1 if record.status == "ACTIVE" else 0

        value = record.effective_date
        if isinstance(value, datetime):
            effective = value
        elif isinstance(value, date):
            effective = datetime.combine(value, datetime.min.time())
        else:
            effective = datetime.min

        return active_rank, effective

    return max(records, key=sort_key)


def build_qualification_lookup(
    population_ssos: set[str],
) -> tuple[dict[tuple[str, str], QualificationRecord], list[str]]:
    settings = load_settings()
    qualification_folder = Path(settings["qualification_folder"])

    grouped: dict[tuple[str, str], list[QualificationRecord]] = {}
    warnings: list[str] = []

    for module_code, filename in QUALIFICATION_FILES.items():
        file_path = qualification_folder / filename
        if not file_path.exists():
            raise FileNotFoundError(f"Qualification file not found: {file_path}")

        records = read_qualification_file(
            file_path,
            module_code,
            population_ssos,
        )

        for record in records:
            grouped.setdefault((record.sso, record.module_code), []).append(record)

    lookup: dict[tuple[str, str], QualificationRecord] = {}

    for key, records in grouped.items():
        selected = choose_record(records)
        lookup[key] = selected

        if len(records) > 1:
            warnings.append(
                f"Duplicate qualification rows: {selected.sso} | "
                f"{selected.module_code} | selected newest preferred record"
            )

        if selected.status != "ACTIVE":
            warnings.append(
                f"Inactive qualification for roster employee: {selected.sso} | "
                f"{selected.module_code} | status {selected.status or 'BLANK'}"
            )

    return lookup, warnings


def format_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    text = clean_text(value)
    if not text:
        return ""

    for date_format in ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, date_format).strftime("%Y-%m-%d")
        except ValueError:
            continue

    return text


def main() -> int:
    people = get_report_population()
    population_ssos = {person.sso for person in people}
    people_by_sso = {person.sso: person for person in people}

    lookup, warnings = build_qualification_lookup(population_ssos)

    print("=" * 76)
    print("PQS DARK UNIVERSE REPORT - QUALIFICATION SUMMARY")
    print("=" * 76)
    print(f"Roster employees: {len(population_ssos)}")
    print(f"Qualification matches: {len(lookup)}")

    for module_code in QUALIFICATION_FILES:
        records = [
            record
            for (sso, module), record in lookup.items()
            if module == module_code
        ]
        active_count = sum(record.status == "ACTIVE" for record in records)
        inactive_count = len(records) - active_count

        print(
            f"{module_code}: {len(records)} roster match(es) | "
            f"ACTIVE {active_count} | OTHER {inactive_count}"
        )

    print("\nMATCHED QUALIFICATIONS")
    print("-" * 76)

    for module_code in QUALIFICATION_FILES:
        module_records = sorted(
            (
                record
                for (sso, module), record in lookup.items()
                if module == module_code
            ),
            key=lambda record: (
                people_by_sso[record.sso].population,
                people_by_sso[record.sso].roster_order,
            ),
        )

        if not module_records:
            continue

        print(f"\n{module_code}")
        for record in module_records:
            person = people_by_sso[record.sso]
            print(
                f"  {record.sso} | {person.name} | {person.population} | "
                f"{record.status or 'BLANK'} | {format_date(record.effective_date)}"
            )

    print("\nWARNINGS")
    print("-" * 76)

    if warnings:
        for warning in warnings:
            print(f"  {warning}")
    else:
        print("  No qualification warnings found for current roster employees.")

    inactive_keys = [
        key for key, record in lookup.items() if record.status != "ACTIVE"
    ]

    if inactive_keys:
        print(
            f"\nQUALIFICATION VALIDATION COMPLETED WITH "
            f"{len(inactive_keys)} INACTIVE/OTHER STATUS RECORD(S)."
        )
    else:
        print("\nQUALIFICATION VALIDATION PASSED.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
