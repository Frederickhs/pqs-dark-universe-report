from __future__ import annotations

import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from employee_summary_processor import (
    ACTION_PRIORITY,
    build_employee_summaries,
    venue_status,
    venue_work_orders,
)
from module_result_processor import load_settings, required_modules_for
from snapshot_manager import build_results_with_snapshot, write_snapshot


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FOLDER = PROJECT_ROOT / "Output"
ARCHIVE_OUTPUT_FOLDER = PROJECT_ROOT / "Archive" / "Outputs"
OUTPUT_FILE = OUTPUT_FOLDER / "PQS Dark Universe Report.xlsx"

POPULATION_ORDER = [
    "1st Shift",
    "2nd Shift",
    "3rd Shift Wolf",
    "3rd Shift Unchained",
    "Supervisors",
    "Animation",
]

VENUE_ORDER = ["Wolf", "Darkmoor", "Unchained", "Animation"]

DARK_BLUE = "17365D"
MEDIUM_BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6E0B4"
DARK_GREEN = "548235"
YELLOW = "FFF2CC"
ORANGE = "F4B183"
RED = "F4CCCC"
DARK_RED = "C00000"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
PURPLE = "E4DFEC"
TEAL = "DDEBF7"
BLACK = "000000"

THIN_GRAY = Side(style="thin", color="B7B7B7")
BOTTOM_BORDER = Border(bottom=THIN_GRAY)

STATUS_FILLS = {
    "Completed": GREEN,
    "Not Required": GRAY,
    "Completed - Awaiting Maximo Status": PURPLE,
    "Review Required": YELLOW,
    "In Progress": LIGHT_BLUE,
    "No Progress": ORANGE,
    "Not Started": GRAY,
    "No Maximo Record": RED,
    "ACK Error - Technician Action Required": RED,
    "Task Total Changed - Review Data": RED,
    "Progress Decreased - Review Data": RED,
}


def archive_existing_output(run_timestamp: datetime) -> Path | None:
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    existing = [path for path in OUTPUT_FOLDER.iterdir() if path.is_file()]

    if not existing:
        return None

    archive_folder = ARCHIVE_OUTPUT_FOLDER / run_timestamp.strftime("%Y-%m-%d_%H%M%S")
    archive_folder.mkdir(parents=True, exist_ok=False)

    for path in existing:
        shutil.move(str(path), str(archive_folder / path.name))

    return archive_folder


def progress_text(result) -> str:
    if result.completed_tasks is None and result.total_tasks is None:
        return "Not Required" if result.result == "Not Required" else ""

    completed = result.completed_tasks if result.completed_tasks is not None else 0
    total = result.total_tasks if result.total_tasks is not None else ""

    if isinstance(completed, float) and completed.is_integer():
        completed = int(completed)
    if isinstance(total, float) and total.is_integer():
        total = int(total)

    return f"{completed}/{total}"


def percentage_value(result):
    if result.result == "Not Required":
        return "Not Required"
    return result.percentage if result.percentage is not None else ""


def module_venue(module_code: str) -> str:
    if module_code.startswith("CU"):
        return "Wolf"
    if module_code.startswith("DU"):
        return "Darkmoor"
    return "Unchained"


def population_modules(population: str, people, settings) -> list[str]:
    rows = [person for person in people if person.population == population]
    modules: list[str] = []

    for person in rows:
        for module in required_modules_for(person, settings):
            if module not in modules:
                modules.append(module)

    return modules


def population_venues(population: str, people) -> list[str]:
    if population == "Animation":
        return ["Animation"]
    if population in {"1st Shift", "2nd Shift"}:
        return ["Wolf", "Darkmoor", "Unchained"]
    if population == "3rd Shift Wolf":
        return ["Wolf", "Darkmoor"]
    if population == "3rd Shift Unchained":
        return ["Unchained"]
    if population == "Supervisors":
        return ["Wolf", "Darkmoor", "Unchained"]
    return []


def style_title(ws, title: str, last_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26


def style_headers(ws, row: int, last_column: int) -> None:
    for cell in ws[row][:last_column]:
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOTTOM_BORDER
    ws.row_dimensions[row].height = 36


def apply_status_fill(cell) -> None:
    value = str(cell.value or "")
    fill = STATUS_FILLS.get(value)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
        if value.startswith("ACK") or value.startswith("Progress Decreased") or value.startswith("Task Total"):
            cell.font = Font(color=DARK_RED, bold=True)


def set_widths(ws, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        if header == "Name":
            width = 26
        elif header in {"SSO", "Position", "Venue", "Shift"}:
            width = 15
        elif header.endswith("Progress"):
            width = 13
        elif header.endswith("%"):
            width = 11
        elif "Work Order" in header:
            width = 24
        elif header.endswith("PQS") or header == "PQS Status":
            width = 20
        elif header == "Missing PQS":
            width = 30
        elif header == "Notes":
            width = 55
        else:
            width = 24
        ws.column_dimensions[get_column_letter(index)].width = width


def create_population_sheet(wb, population, people, results, summaries, settings) -> None:
    ws = wb.create_sheet(population)
    population_people = [person for person in people if person.population == population]
    summary_lookup = {summary.sso: summary for summary in summaries}
    modules = population_modules(population, people, settings)
    venues = population_venues(population, people)

    headers = ["Name", "SSO", "Position", "Venue", "Shift"]

    for module in modules:
        headers.extend([module, f"{module} Progress", f"{module} %"])

    for venue in venues:
        headers.extend([f"{venue} Work Order(s)", f"{venue} Full PQS"])

    headers.extend(["PQS Status", "Missing PQS", "Notes"])

    style_title(ws, f"PQS Dark Universe Report - {population}", len(headers))
    ws.append(headers)
    style_headers(ws, 2, len(headers))

    module_status_columns: list[int] = []
    percentage_columns: list[int] = []
    venue_status_columns: list[int] = []

    for person in population_people:
        summary = summary_lookup[person.sso]
        row = [person.name, person.sso, person.position, person.venue, person.shift]

        for module in modules:
            key = (person.sso, module)
            if key in results:
                result = results[key]
                row.extend([result.result, progress_text(result), percentage_value(result)])
            else:
                row.extend(["Not Required", "Not Required", "Not Required"])

        for venue in venues:
            row.extend([venue_work_orders(summary, venue), venue_status(summary, venue) or "Not Required"])

        row.extend([summary.pqs_status, summary.missing_pqs, summary.notes])
        ws.append(row)

    for column, header in enumerate(headers, start=1):
        if header in modules:
            module_status_columns.append(column)
        elif header.endswith(" %"):
            percentage_columns.append(column)
        elif header.endswith("Full PQS"):
            venue_status_columns.append(column)

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BOTTOM_BORDER

        ws.cell(row[0].row, 2).number_format = "@"

        for column in module_status_columns:
            apply_status_fill(ws.cell(row[0].row, column))

        for column in percentage_columns:
            cell = ws.cell(row[0].row, column)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"

        for column in venue_status_columns:
            cell = ws.cell(row[0].row, column)
            if cell.value == "COMPLETE":
                cell.fill = PatternFill("solid", fgColor=GREEN)
            elif cell.value == "INCOMPLETE":
                cell.fill = PatternFill("solid", fgColor=RED)

        pqs_column = headers.index("PQS Status") + 1
        pqs_cell = ws.cell(row[0].row, pqs_column)
        pqs_cell.fill = PatternFill(
            "solid", fgColor=GREEN if pqs_cell.value == "Full PQS" else RED
        )
        pqs_cell.font = Font(bold=True)

    # Hide a module's Progress and % columns when that sheet has no actual
    # Maximo progress values for the module. Keep both visible if at least
    # one employee has task progress or a percentage.
    for module in modules:
        progress_column = headers.index(f"{module} Progress") + 1
        percentage_column = headers.index(f"{module} %") + 1

        has_progress_data = any(
            results.get((person.sso, module)) is not None
            and (
                results[(person.sso, module)].completed_tasks is not None
                or results[(person.sso, module)].total_tasks is not None
                or results[(person.sso, module)].percentage is not None
            )
            for person in population_people
        )

        if not has_progress_data:
            ws.column_dimensions[get_column_letter(progress_column)].hidden = True
            ws.column_dimensions[get_column_letter(percentage_column)].hidden = True

    # Hide a venue Work Order(s) column when every employee on the sheet
    # has a blank work-order value. The venue Full PQS column stays visible.
    for venue in venues:
        work_order_header = f"{venue} Work Order(s)"
        work_order_column = headers.index(work_order_header) + 1

        has_work_order_data = any(
            bool(venue_work_orders(summary_lookup[person.sso], venue).strip())
            for person in population_people
        )

        if not has_work_order_data:
            ws.column_dimensions[get_column_letter(work_order_column)].hidden = True

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{ws.max_row}"
    ws.sheet_view.showGridLines = False
    set_widths(ws, headers)


def add_metric_table(ws, start_row, title, rows, include_no_maximo):
    headers = ["Group", "Total", "Complete", "Incomplete"]
    if include_no_maximo:
        headers.append("No Maximo Record")
    headers.append("Completion %")

    display_rows = []
    for values in rows:
        if include_no_maximo:
            display_rows.append(values)
        else:
            display_rows.append(values[:4] + values[5:])

    ws.cell(start_row, 1, title)
    ws.cell(start_row, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.cell(start_row, 1).font = Font(color=WHITE, bold=True, size=12)
    ws.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=len(headers),
    )

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row + 1, col, header)
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")

    percentage_column = len(headers)

    for offset, values in enumerate(display_rows, start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(start_row + offset, col, value)
            cell.border = BOTTOM_BORDER
            if col == percentage_column and isinstance(value, (int, float)):
                cell.number_format = "0.0%"

    return start_row + 1, start_row + 1 + len(display_rows)


def create_dashboard(wb, summaries, results) -> None:
    ws = wb.create_sheet("Dash", 0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:N1")
    ws["A1"] = "PQS Dark Universe Dashboard"
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws["A1"].font = Font(color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.row_dimensions[1].height = 30

    population_rows = []
    for population in POPULATION_ORDER:
        group = [summary for summary in summaries if summary.population == population]
        complete = sum(summary.pqs_status == "Full PQS" for summary in group)
        no_maximo = sum(summary.action_category == "No Maximo Record" for summary in group)
        population_rows.append([
            population,
            len(group),
            complete,
            len(group) - complete,
            no_maximo if no_maximo else "",
            complete / len(group) if group else 0,
        ])

    excluding_animation = [s for s in summaries if s.population != "Animation"]
    including_animation = list(summaries)

    for label, group in [
        ("Overall excluding Animation", excluding_animation),
        ("Overall including Animation", including_animation),
    ]:
        complete = sum(summary.pqs_status == "Full PQS" for summary in group)
        no_maximo = sum(summary.action_category == "No Maximo Record" for summary in group)
        population_rows.append([
            label,
            len(group),
            complete,
            len(group) - complete,
            no_maximo if no_maximo else "",
            complete / len(group) if group else 0,
        ])

    include_no_maximo = any(
        summary.action_category == "No Maximo Record"
        for summary in summaries
    )

    population_header_row, population_end_row = add_metric_table(
        ws,
        3,
        "Total Completed",
        population_rows,
        include_no_maximo,
    )

    venue_rows = []
    for venue in VENUE_ORDER:
        applicable = [
            summary
            for summary in summaries
            if any(item.venue == venue for item in summary.venue_summaries)
        ]
        complete = sum(venue_status(summary, venue) == "COMPLETE" for summary in applicable)
        no_maximo = sum(summary.action_category == "No Maximo Record" for summary in applicable)
        venue_rows.append([
            venue,
            len(applicable),
            complete,
            len(applicable) - complete,
            no_maximo if no_maximo else "",
            complete / len(applicable) if applicable else 0,
        ])

    venue_start = population_end_row + 3
    venue_header_row, venue_end_row = add_metric_table(
        ws,
        venue_start,
        "Completion by Venue",
        venue_rows,
        include_no_maximo,
    )

    review_employees = sum(bool(summary.review_modules) for summary in summaries)
    review_modules = sum(len(summary.review_modules) for summary in summaries)

    kpi_row = venue_end_row + 3
    ws.cell(kpi_row, 1, "Review Metrics")
    ws.cell(kpi_row, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.cell(kpi_row, 1).font = Font(color=WHITE, bold=True, size=12)
    ws.merge_cells(start_row=kpi_row, start_column=1, end_row=kpi_row, end_column=4)

    ws.cell(kpi_row + 1, 1, "Employees Requiring Review")
    ws.cell(kpi_row + 1, 2, review_employees)
    ws.cell(kpi_row + 2, 1, "Modules Requiring Review")
    ws.cell(kpi_row + 2, 2, review_modules)

    action_counter = Counter(summary.action_category for summary in summaries)
    action_row = kpi_row + 4
    ws.cell(action_row, 1, "Action Needed")
    ws.cell(action_row, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.cell(action_row, 1).font = Font(color=WHITE, bold=True, size=12)
    ws.merge_cells(start_row=action_row, start_column=1, end_row=action_row, end_column=2)
    ws.cell(action_row + 1, 1, "Category")
    ws.cell(action_row + 1, 2, "Employees")

    action_categories = [category for category in ACTION_PRIORITY if action_counter.get(category)]
    for index, category in enumerate(action_categories, start=2):
        ws.cell(action_row + index, 1, category)
        ws.cell(action_row + index, 2, action_counter[category])

    action_end_row = action_row + 1 + len(action_categories)

    list_row = action_end_row + 3
    ws.cell(list_row, 1, "Employee Action List")
    ws.cell(list_row, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.cell(list_row, 1).font = Font(color=WHITE, bold=True, size=12)
    ws.merge_cells(start_row=list_row, start_column=1, end_row=list_row, end_column=7)

    action_headers = ["Category", "Name", "SSO", "Population", "Venue", "Missing PQS", "Notes"]
    for col, header in enumerate(action_headers, start=1):
        cell = ws.cell(list_row + 1, col, header)
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.font = Font(color=WHITE, bold=True)

    actionable = [summary for summary in summaries if summary.action_category != "Completed"]
    actionable.sort(key=lambda item: (ACTION_PRIORITY.index(item.action_category), item.population, item.roster_order))

    for row_index, summary in enumerate(actionable, start=list_row + 2):
        values = [
            summary.action_category,
            summary.name,
            summary.sso,
            summary.population,
            summary.venue,
            summary.missing_pqs,
            summary.notes,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BOTTOM_BORDER
        ws.cell(row_index, 3).number_format = "@"

    population_chart = BarChart()
    population_chart.type = "col"
    population_chart.style = 10
    population_chart.title = "Total Completed"
    population_chart.y_axis.title = "Employees"
    population_chart.x_axis.title = "Population"
    data = Reference(ws, min_col=3, max_col=4, min_row=population_header_row, max_row=population_end_row)
    cats = Reference(ws, min_col=1, min_row=population_header_row + 1, max_row=population_end_row)
    population_chart.add_data(data, titles_from_data=True)
    population_chart.set_categories(cats)
    population_chart.height = 8
    population_chart.width = 14
    population_chart.legend.position = "b"
    ws.add_chart(population_chart, "H3")

    venue_chart = BarChart()
    venue_chart.type = "col"
    venue_chart.style = 10
    venue_chart.title = "Completion by Venue"
    venue_chart.y_axis.title = "Employees"
    venue_chart.x_axis.title = "Venue"
    data = Reference(ws, min_col=3, max_col=4, min_row=venue_header_row, max_row=venue_end_row)
    cats = Reference(ws, min_col=1, min_row=venue_header_row + 1, max_row=venue_end_row)
    venue_chart.add_data(data, titles_from_data=True)
    venue_chart.set_categories(cats)
    venue_chart.height = 8
    venue_chart.width = 14
    venue_chart.legend.position = "b"
    ws.add_chart(venue_chart, "H20")

    action_chart = BarChart()
    action_chart.type = "bar"
    action_chart.style = 10
    action_chart.title = "Action Needed"
    action_chart.x_axis.title = "Employees"
    data = Reference(ws, min_col=2, min_row=action_row + 1, max_row=action_end_row)
    cats = Reference(ws, min_col=1, min_row=action_row + 2, max_row=action_end_row)
    action_chart.add_data(data, titles_from_data=True)
    action_chart.set_categories(cats)
    action_chart.height = 8
    action_chart.width = 14
    action_chart.legend = None
    ws.add_chart(action_chart, "H37")

    widths = {1: 34, 2: 24, 3: 14, 4: 24, 5: 18, 6: 32, 7: 70}
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width

    ws.freeze_panes = "A3"


def validate_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    expected = ["Dash"] + POPULATION_ORDER

    if workbook.sheetnames != expected:
        raise ValueError(
            f"Unexpected sheet order. Expected {expected}, got {workbook.sheetnames}"
        )

    for sheet_name in expected:
        worksheet = workbook[sheet_name]
        if worksheet.max_row < 2 or worksheet.max_column < 2:
            raise ValueError(f"Generated sheet appears empty: {sheet_name}")

    workbook.close()


def generate_report() -> tuple[Path, Path | None, Path]:
    run_timestamp = datetime.now()

    (
        people,
        results,
        warnings,
        unknown_modules,
        _previous_path,
        _snapshot_path,
        _comparison_notes,
    ) = build_results_with_snapshot(
        run_timestamp=run_timestamp,
        save_snapshot=False,
    )

    if warnings:
        raise ValueError("Processing warnings block report generation: " + "; ".join(warnings))
    if unknown_modules:
        raise ValueError("Unknown module descriptions block report generation")

    summaries = build_employee_summaries(people, results)
    settings = load_settings()

    archived_folder = archive_existing_output(run_timestamp)

    workbook = Workbook()
    workbook.remove(workbook.active)

    create_dashboard(workbook, summaries, results)

    for population in POPULATION_ORDER:
        create_population_sheet(
            workbook,
            population,
            people,
            results,
            summaries,
            settings,
        )

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_FILE)
    validate_workbook(OUTPUT_FILE)

    snapshot_path = write_snapshot(results, run_timestamp)
    return OUTPUT_FILE, archived_folder, snapshot_path


def main() -> int:
    output_file, archived_folder, snapshot_path = generate_report()

    print("=" * 82)
    print("PQS DARK UNIVERSE REPORT - EXCEL GENERATION")
    print("=" * 82)
    print(f"Output workbook: {output_file}")
    print(
        "Archived previous output: "
        + (str(archived_folder) if archived_folder else "No previous output files")
    )
    print(f"Successful snapshot: {snapshot_path}")
    print("Visible sheets:")
    print("  Dash")
    for population in POPULATION_ORDER:
        print(f"  {population}")
    print("\nREPORT GENERATION PASSED.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
