from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

LINE_SCHEDULE_FILE = PROJECT_ROOT / "Input" / "Line Schedule" / "Line Schedule.xlsx"
ANIMATION_ROSTER_FILE = PROJECT_ROOT / "Input" / "Rosters" / "Animation Roster.xlsx"
SUPERVISOR_ROSTER_FILE = PROJECT_ROOT / "Input" / "Supervisors" / "Supervisor Roster.xlsx"


@dataclass
class RosterPerson:
    name: str
    sso: str
    position: str
    shift: str
    venue: str
    population: str
    roster_order: int


LINE_SCHEDULE_RULES = {
    "1st Shift": {
        "shift": "1st",
        "venue": "Dark Universe",
        "population": "1st Shift",
    },
    "2nd Shift": {
        "shift": "2nd",
        "venue": "Dark Universe",
        "population": "2nd Shift",
    },
    "3rd Shift Wolf": {
        "shift": "3rd",
        "venue": "Wolf",
        "population": "3rd Shift Wolf",
    },
    "3rd Shift Unchained": {
        "shift": "3rd",
        "venue": "Unchained",
        "population": "3rd Shift Unchained",
    },
}

POSITION_LABELS = {
    "lead": "Lead",
    "show controls": "Show Controls",
    "show control": "Show Controls",
    "mechanics": "Mechanics",
    "mechanic": "Mechanics",
}

IGNORE_NAME_PREFIXES = (
    "open",
    "daily total",
    "nightly total",
)

IGNORE_EXACT_VALUES = {
    "technician",
    "supervisor",
    "apprentice",
}


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def normalize_sso(value: object) -> str:
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = clean_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]

    return text


def is_ignored_name(name: str) -> bool:
    normalized = name.casefold()
    if normalized in IGNORE_EXACT_VALUES:
        return True

    return any(normalized.startswith(prefix) for prefix in IGNORE_NAME_PREFIXES)


def read_line_schedule() -> list[RosterPerson]:
    workbook = load_workbook(LINE_SCHEDULE_FILE, read_only=True, data_only=True)
    people: list[RosterPerson] = []

    for sheet_name, rules in LINE_SCHEDULE_RULES.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Line Schedule worksheet not found: {sheet_name}")

        worksheet = workbook[sheet_name]
        worksheet.reset_dimensions()
        current_position = ""
        roster_order = 0

        for row in worksheet.iter_rows(values_only=True):
            values = [clean_text(value) for value in row]
            if not any(values):
                continue

            first_value = values[0]
            if not first_value:
                continue

            normalized_first = first_value.casefold()
            if normalized_first in POSITION_LABELS:
                current_position = POSITION_LABELS[normalized_first]
                continue

            if is_ignored_name(first_value):
                continue

            sso = ""
            for value in row:
                candidate = normalize_sso(value)
                if candidate.isdigit() and len(candidate) == 9:
                    sso = candidate
                    break

            if not sso:
                continue

            roster_order += 1
            people.append(
                RosterPerson(
                    name=first_value.replace("*", "").strip(),
                    sso=sso,
                    position=current_position or "Unknown",
                    shift=rules["shift"],
                    venue=rules["venue"],
                    population=rules["population"],
                    roster_order=roster_order,
                )
            )

    workbook.close()
    return people


def read_animation_roster() -> list[RosterPerson]:
    workbook = load_workbook(ANIMATION_ROSTER_FILE, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    worksheet.reset_dimensions()
    people: list[RosterPerson] = []
    roster_order = 0

    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number == 1:
            continue

        name = clean_text(row[0] if len(row) > 0 else None)
        sso = normalize_sso(row[1] if len(row) > 1 else None)
        if not name or not sso:
            continue

        roster_order += 1
        people.append(
            RosterPerson(
                name=name,
                sso=sso,
                position="Show Control",
                shift="3rd",
                venue="Animation",
                population="Animation",
                roster_order=roster_order,
            )
        )

    workbook.close()
    return people


def read_supervisor_roster() -> list[RosterPerson]:
    workbook = load_workbook(SUPERVISOR_ROSTER_FILE, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    worksheet.reset_dimensions()

    header_row = next(worksheet.iter_rows(min_row=1, max_row=1))
    headers = {
        clean_text(cell.value).casefold(): cell.column
        for cell in header_row
        if clean_text(cell.value)
    }

    required_headers = {"name", "sso", "shift", "venue"}
    missing_headers = required_headers - set(headers)
    if missing_headers:
        raise ValueError(
            "Supervisor roster missing columns: " + ", ".join(sorted(missing_headers))
        )

    people: list[RosterPerson] = []
    roster_order = 0

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        name = clean_text(row[headers["name"] - 1])
        sso = normalize_sso(row[headers["sso"] - 1])
        shift = clean_text(row[headers["shift"] - 1])
        venue = clean_text(row[headers["venue"] - 1])

        if not name and not sso:
            continue
        if not sso:
            raise ValueError(f"Supervisor is missing SSO: {name or 'Unknown name'}")

        if shift in {"1st", "2nd"}:
            venue = "Dark Universe"
        elif shift == "3rd" and venue not in {"Wolf", "Unchained"}:
            raise ValueError(f"Invalid 3rd Shift supervisor venue: {name} | {venue}")

        roster_order += 1
        people.append(
            RosterPerson(
                name=name,
                sso=sso,
                position="Supervisor",
                shift=shift,
                venue=venue,
                population="Supervisors",
                roster_order=roster_order,
            )
        )

    workbook.close()
    return people


def remove_special_population_duplicates(
    line_schedule_people: list[RosterPerson],
    animation_people: list[RosterPerson],
    supervisor_people: list[RosterPerson],
) -> tuple[list[RosterPerson], list[str]]:
    animation_ssos = {person.sso for person in animation_people}
    supervisor_ssos = {person.sso for person in supervisor_people}
    special_ssos = animation_ssos | supervisor_ssos
    messages: list[str] = []
    filtered_line_schedule: list[RosterPerson] = []

    for person in line_schedule_people:
        if person.sso in special_ssos:
            destinations = []
            if person.sso in animation_ssos:
                destinations.append("Animation")
            if person.sso in supervisor_ssos:
                destinations.append("Supervisors")

            messages.append(
                f"{person.sso} | {person.name} removed from "
                f"{person.population}; assigned to {', '.join(destinations)}"
            )
        else:
            filtered_line_schedule.append(person)

    overlap = animation_ssos & supervisor_ssos
    for sso in sorted(overlap):
        messages.append(f"ERROR: {sso} appears in both Animation and Supervisors")

    return filtered_line_schedule, messages


def find_duplicate_population_ssos(people: list[RosterPerson]) -> list[str]:
    assignments: dict[str, list[str]] = {}

    for person in people:
        assignments.setdefault(person.sso, []).append(person.population)

    return [
        f"{sso}: {', '.join(populations)}"
        for sso, populations in assignments.items()
        if len(populations) > 1
    ]


def print_population_summary(people: list[RosterPerson]) -> None:
    populations: dict[str, list[RosterPerson]] = {}

    for person in people:
        populations.setdefault(person.population, []).append(person)

    print("=" * 72)
    print("PQS DARK UNIVERSE REPORT - ROSTER SUMMARY")
    print("=" * 72)

    for population in [
        "1st Shift",
        "2nd Shift",
        "3rd Shift Wolf",
        "3rd Shift Unchained",
        "Supervisors",
        "Animation",
    ]:
        population_people = populations.get(population, [])
        print(f"\n{population}: {len(population_people)}")

        for person in population_people:
            print(
                f"  {person.roster_order:>2}. {person.sso} | {person.name} | "
                f"{person.position} | {person.venue}"
            )


def main() -> int:
    line_schedule_people = read_line_schedule()
    animation_people = read_animation_roster()
    supervisor_people = read_supervisor_roster()

    filtered_schedule, messages = remove_special_population_duplicates(
        line_schedule_people,
        animation_people,
        supervisor_people,
    )

    all_people = filtered_schedule + supervisor_people + animation_people
    duplicate_assignments = find_duplicate_population_ssos(all_people)

    print_population_summary(all_people)

    print("\n" + "=" * 72)
    print("ROSTER ASSIGNMENT MESSAGES")
    print("=" * 72)

    if messages:
        for message in messages:
            print(f"  {message}")
    else:
        print("  No employees required reassignment.")

    if duplicate_assignments:
        print("\nERROR: Duplicate final population assignments found:")
        for duplicate in duplicate_assignments:
            print(f"  {duplicate}")
        return 1

    print(f"\nROSTER VALIDATION PASSED: {len(all_people)} unique employees.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
